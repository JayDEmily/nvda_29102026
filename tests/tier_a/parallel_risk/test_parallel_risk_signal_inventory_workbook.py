from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

WORKBOOK = Path(
    "data/reference/signal_workbooks/2026-03-25_NVDA_SIGNAL_WORKBOOK_v3_bounds_handoff_copy.xlsx"
)
SHEET = "Parallel Risk Lane Signal Inventory"
REQUIRED_HEADERS = {
    "signal_id",
    "signal_family",
    "inventory_status",
    "raw_workbook_lineage_ref",
    "derived_workbook_lineage_ref",
    "runtime_contract_target",
    "lawful_prl_use",
    "missing_handling",
    "unavailable_handling",
    "exclusion_reason",
    "notes",
}
SUPPORTED_SIGNALS = {
    "source_time_utc",
    "scheduled_event_window_active",
    "vix_level",
    "spot_price",
    "front_dte",
    "next_dte",
}
DEFERRED_OR_MISSING_RICH_SIGNALS = {
    "distance_to_vwap",
    "dense_options_chain",
    "gamma_exposure_flip",
    "headline_risk_context",
    "crowding_real_money_flow",
}
FORBIDDEN_ADVICE_TERMS = {
    "buy",
    "sell",
    "fade",
    "chase",
    "hedge",
    "size up",
    "stand down",
    "take the trade",
    "avoid the trade",
}
RUNTIME_PRL_PATHS = [
    Path("src/nvda_desk/contracts/parallel_risk.py"),
    *sorted(Path("src/nvda_desk/domain/parallel_risk").glob("*.py")),
    Path("src/nvda_desk/risk/parallel/evaluate.py"),
    Path("src/nvda_desk/services/parallel_risk_lane_service.py"),
]


def _inventory_rows() -> dict[str, dict[str, str]]:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    assert SHEET in workbook.sheetnames
    sheet = workbook[SHEET]
    headers = [cell.value for cell in sheet[3]]
    rows: dict[str, dict[str, str]] = {}
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if not row or row[0] is None:
            continue
        record = {str(header): "" if value is None else str(value) for header, value in zip(headers, row)}
        rows[record["signal_id"]] = record
    return rows


def test_parallel_risk_signal_inventory_workbook_sheet_and_headers_exist() -> None:
    assert WORKBOOK.exists()
    rows = _inventory_rows()
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    headers = {cell.value for cell in workbook[SHEET][3]}

    assert REQUIRED_HEADERS <= headers
    assert len(rows) >= len(SUPPORTED_SIGNALS | DEFERRED_OR_MISSING_RICH_SIGNALS)


def test_parallel_risk_signal_inventory_represents_supported_runtime_signals() -> None:
    rows = _inventory_rows()

    assert SUPPORTED_SIGNALS <= set(rows)
    for signal_id in SUPPORTED_SIGNALS:
        record = rows[signal_id]
        assert record["inventory_status"] == "available_now"
        assert record["required_for_current_atom"] == "yes"
        assert record["runtime_use_permission"] == "runtime_allowed_from_typed_inputs_only"
        assert record["lawful_prl_use"] == "yes_after_typed_promotion"
        assert record["raw_workbook_lineage_ref"] or record["derived_workbook_lineage_ref"]


def test_parallel_risk_signal_inventory_marks_rich_signals_unavailable_not_fabricated() -> None:
    rows = _inventory_rows()

    assert DEFERRED_OR_MISSING_RICH_SIGNALS <= set(rows)
    for signal_id in DEFERRED_OR_MISSING_RICH_SIGNALS:
        record = rows[signal_id]
        assert record["inventory_status"] in {"deferred", "missing"}
        assert record["lawful_prl_use"] == "no_until_typed_source_exists"
        assert record["runtime_use_permission"] == "not_runtime_allowed"
        assert "no_value" in record["unavailable_handling"]
        assert "fabricated" in record["notes"].lower() or "lawful source" in record["notes"].lower()


def test_parallel_risk_signal_inventory_excluded_authority_is_not_runtime_demanded() -> None:
    rows = _inventory_rows()

    for signal_id in {"capital_sizing_signal", "order_instruction_signal"}:
        record = rows[signal_id]
        assert record["inventory_status"] == "excluded"
        assert record["required_for_current_atom"] == "no"
        assert record["unavailable_handling"] == "excluded_signal_not_demanded"
        assert record["exclusion_reason"]


def test_parallel_risk_signal_inventory_and_runtime_boundary_have_no_advice_or_workbook_reads() -> None:
    rows = _inventory_rows()
    workbook_text = "\n".join(
        "\n".join(record.values()).lower() for record in rows.values()
    )
    for forbidden in FORBIDDEN_ADVICE_TERMS:
        assert forbidden not in workbook_text

    runtime_text = "\n".join(path.read_text(encoding="utf-8") for path in RUNTIME_PRL_PATHS)
    forbidden_runtime_tokens = (".xlsx", "openpyxl", "load_workbook", "SpreadsheetFile", "import_xlsx")
    for token in forbidden_runtime_tokens:
        assert token not in runtime_text
